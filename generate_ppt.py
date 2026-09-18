from pptx import Presentation
from pptx.util import Inches, Pt

def add_slide(prs, title, content_bullets):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title_placeholder = slide.shapes.title
    body_placeholder = slide.shapes.placeholders[1]

    title_placeholder.text = title
    tf = body_placeholder.text_frame
    tf.clear() # Clear existing text
    
    for bullet in content_bullets:
        p = tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(18)

prs = Presentation()

slides_data = [
    {
        "title": "1. ABSTRACT",
        "bullets": [
            "Overview: AutoLingo is an AI-powered smart text autocomplete engine designed to predict and generate text sequentially.",
            "Core Technology: Built on a PyTorch Embedding-to-Linear Predictor, mapping character contexts to next-character probabilities.",
            "Data Pipeline: Uses semi-supervised learning (Label Propagation, Tri-Training) to automatically sanitize and filter noisy training data.",
            "Comparative Analysis: Evaluates the neural network against traditional ML baselines (KNN, Decision Trees, Logistic Regression).",
            "Forecasting: Implements regression models to predict hyperparameter impacts on validation loss and training time.",
            "Deployment: Features a live, interactive web dashboard using FastAPI and Streamlit/Dash for real-time autocomplete suggestions."
        ]
    },
    {
        "title": "2. INTRODUCTION",
        "bullets": [
            "The Problem: Modern applications require efficient, context-aware autocomplete systems to improve user experience and typing speed.",
            "Challenges: Training language models requires clean data, but real-world text (e.g., web scraping) is often noisy and unstructured. Manual labeling is expensive.",
            "The Solution (AutoLingo): An end-to-end pipeline that handles everything from data ingestion and automated cleaning to model training, benchmarking, and real-time inference.",
            "Objective: To build a scalable autocomplete engine and rigorously compare deep learning approaches against classical machine learning methodologies."
        ]
    },
    {
        "title": "3. LITERATURE SURVEY",
        "bullets": [
            "Traditional Text Prediction: Early systems relied on N-gram models and statistical frequencies, which struggle with long-range context.",
            "Classical ML in NLP: Algorithms like KNN and Decision Trees can be adapted for character/word classification but often face the 'curse of dimensionality' with text data.",
            "Deep Learning Advances: Dense embeddings and neural networks (like PyTorch-based linear models and Transformers) have revolutionized text generation by capturing semantic relationships.",
            "Semi-Supervised Learning: Techniques like Label Propagation and Tri-Training leverage small amounts of labeled data to automatically annotate large, unlabeled datasets."
        ]
    },
    {
        "title": "4. RESEARCH GAPS",
        "bullets": [
            "Data Bottlenecks: Lack of accessible, automated semi-supervised data cleaning pipelines for building bespoke autocomplete engines without heavy manual annotation.",
            "Holistic Benchmarking: Limited platforms offer direct, side-by-side comparisons of deep learning models versus traditional ML baselines (like Random Forest or KNN) in real-time text generation tasks.",
            "Resource Forecasting: Difficulty in predicting a model's training latency and validation loss based on architectural hyperparameters before committing significant compute resources."
        ]
    },
    {
        "title": "5. ARCHITECTURE OF PROPOSED WORK",
        "bullets": [
            "Data Ingestion & Labeling: Raw text is cleaned, and semi-supervised Label Propagation filters out junk text to create a clean corpus.",
            "Core ML Engines: The cleaned data flows into two parallel training tracks (PyTorch TextPredictor vs. Classical Classifiers).",
            "Benchmarking & Forecasting: Model outputs are evaluated, and regression models (Polynomial/Linear) are fitted to forecast performance.",
            "Interactive Deployment: A FastAPI backend serves the trained models to a live Web UI Dashboard for real-time user interaction."
        ]
    },
    {
        "title": "6. PROPOSED MODEL",
        "bullets": [
            "Data Sanitizer: Character-level n-gram frequency matrices combined with Label Propagation to classify sentences as high-quality or noise.",
            "Neural Network (Core): A PyTorch-based model utilizing position embeddings and hidden linear transformation layers, optimized via Cross-Entropy Loss and AdamW.",
            "Traditional ML Baselines: Context windows encoded as one-hot vectors and fed into scikit-learn models (KNN, Decision Tree, Logistic Regression).",
            "Performance Regressor: Linear and Polynomial Regression models trained on system logs to predict execution time based on hyperparameters."
        ]
    },
    {
        "title": "7. RESULTS",
        "bullets": [
            "Model Performance: The PyTorch Embedding-to-Linear model demonstrates superior context-awareness and lower Cross-Entropy loss compared to traditional ML baselines.",
            "Data Quality Impact: Tri-Training and Label Propagation successfully increased the signal-to-noise ratio in the training corpus, leading to faster convergence.",
            "Forecasting Accuracy: The Polynomial Regression forecaster accurately predicted training latency for larger model sizes, allowing for optimized hyperparameter tuning.",
            "Real-Time API: The dashboard successfully handles concurrent requests, visually comparing the neural network's predictions against the KNN/Decision Tree outputs in real-time."
        ]
    },
    {
        "title": "8. CONCLUSION",
        "bullets": [
            "Developed a fully integrated, smart text autocomplete system from scratch.",
            "Successfully automated the tedious data cleaning process using semi-supervised learning techniques.",
            "Demonstrated that neural embedding approaches yield much more cohesive autocomplete suggestions than traditional ML models.",
            "Built a robust forecasting and benchmarking suite, culminating in a live, interactive API dashboard that brings the research to life."
        ]
    },
    {
        "title": "9. FUTURE SCOPE",
        "bullets": [
            "Architectural Upgrades: Transitioning the core PyTorch engine to incorporate full Transformer Blocks (Attention mechanisms) for better long-term context.",
            "Tokenization: Shifting from character-level prediction to sub-word (BPE) or word-level tokenization for faster inference and larger vocabulary handling.",
            "Cloud Deployment: Containerizing the API with Docker and deploying to cloud infrastructure (AWS/GCP) for scalability and load balancing.",
            "User Feedback Loop: Expanding the dashboard to include A/B testing, allowing user selections to feed back into and improve the model continuously."
        ]
    },
    {
        "title": "10. REFERENCES",
        "bullets": [
            "PyTorch Documentation for Deep Learning Frameworks.",
            "Scikit-Learn: Machine Learning in Python (Pedregosa et al., 2011).",
            "Xiaojin Zhu and Zoubin Ghahramani. Learning from labeled and unlabeled data with label propagation. (2002).",
            "Zhi-Hua Zhou and Ming Li. Tri-training: Exploiting unlabeled data using three classifiers. (IEEE Transactions on Knowledge and Data Engineering, 2005)."
        ]
    }
]

# Add a title slide
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "AutoLingo: AI-Powered Smart Text Autocomplete Engine"
subtitle.text = "Project Presentation"

# Add content slides
for slide_data in slides_data:
    add_slide(prs, slide_data["title"], slide_data["bullets"])

prs.save('AutoLingo_Presentation.pptx')
print("Successfully generated AutoLingo_Presentation.pptx")
